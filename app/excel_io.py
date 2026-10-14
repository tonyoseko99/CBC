"""Excel export of the class mark sheet. (Uploads use app/templates.py.)"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import grading as G
from .compute import SHEET_DEFAULTS, mark_sheet_layout


# -- export ------------------------------------------------------------------
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
SUM_FILL = PatternFill("solid", fgColor="F2F2F2")
GROUP_FILL = {"EE": "C6EFCE", "ME": "DDEBF7", "AE": "FFEB9C", "BE": "FFC7CE"}


def _style_header(ws, top, n_cols, n_rows=2):
    for r in range(top, top + n_rows):
        for c in range(1, n_cols + 1):
            x = ws.cell(r, c)
            x.font = Font(bold=True, size=9)
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            x.fill = HEAD_FILL
            x.border = BORDER


def sheet_title(res, lay):
    t = res.title()
    if len(lay["rows"]) != len(res.rows):
        t += f"  ({len(lay['rows'])} OF {len(res.rows)} LEARNERS)"
    return t


def export_mark_sheet(path, res, opts=None):
    o = {**SHEET_DEFAULTS, **(opts or {})}
    lay = mark_sheet_layout(res, o)
    n = lay["n_cols"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Mark Sheet"
    s = res.settings
    for i, text in enumerate([s.get("header_line", ""), sheet_title(res, lay)], start=1):
        ws.cell(i, 1, text).font = Font(bold=True, size=12 if i == 2 else 11)
        ws.cell(i, 1).alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=n)
    top = 3
    for ri, row in enumerate(lay["header"]):
        for ci, v in enumerate(row):
            ws.cell(top + ri, ci + 1, v)
    for r0, c0, r1, c1 in lay["spans"]:
        if (r0, c0) != (r1, c1):
            ws.merge_cells(start_row=top + r0, start_column=c0 + 1,
                           end_row=top + r1, end_column=c1 + 1)
    _style_header(ws, top, n)
    r = top + 2
    kinds = lay["kinds"]
    for row in lay["body"]:
        for ci, v in enumerate(row):
            x = ws.cell(r, ci + 1, v if v != "" else None)
            x.border = BORDER
            x.font = Font(size=9, bold=kinds[ci] in ("grd", "level", "pts", "points"))
            if kinds[ci] != "name":
                x.alignment = Alignment(horizontal="center")
            if kinds[ci] in ("grd", "level") and v:
                x.fill = PatternFill("solid", fgColor=GROUP_FILL[str(v)[:2]])
        r += 1
    r += 1
    for row in lay["summary"]:
        for ci, v in enumerate(row):
            x = ws.cell(r, ci + 1, v if v != "" else None)
            x.border = BORDER
            x.fill = SUM_FILL
            x.font = Font(size=9, bold=ci == 1)
            x.alignment = Alignment(horizontal="center" if ci != 1 else "left")
        r += 1
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 28
    for c in range(3, n + 1):
        ws.column_dimensions[get_column_letter(c)].width = 9 if kinds[c - 1] in (
            "total", "outof", "gtotal", "avg", "points", "level") else 6.5
    ws.row_dimensions[top + 1].height = 30
    ws.freeze_panes = ws.cell(top + 2, 3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{top}:{top + 1}"

    if o["analysis"]:
        _summary_sheet(wb.create_sheet("Subject Analysis"), res, lay)
    wb.save(path)


def _summary_sheet(ws, res, lay):
    ws.append([res.settings.get("header_line", "")])
    ws.append([sheet_title(res, lay) + " - SUBJECT ANALYSIS"])
    ws.append([])
    head = ["LEARNING AREA", "SAT", "MEAN %", "MEAN LEVEL"] + G.CODES + G.GROUPS
    ws.append(head)
    _style_header(ws, 4, len(head), 1)
    for st in res.subject_stats(lay["rows"], lay["subjects"]):
        ws.append([st["subject"]["name"], st["sat"],
                   round(st["mean_pct"], 1) if st["mean_pct"] is not None else None,
                   st["band"].code if st["band"] else None]
                  + [st["counts"][c] for c in G.CODES] + [st["groups"][g] for g in G.GROUPS])
    ov = res.overall(lay["rows"])
    ws.append([])
    ws.append(["Learners included", ov["learners"]])
    ws.append(["Learners assessed", ov["with_marks"]])
    if ov["mean_pct"] is not None:
        ws.append(["Class mean %", round(ov["mean_pct"], 1), ov["band"].code])
    for row in ws.iter_rows(min_row=5, max_row=4 + len(lay["subjects"])):
        for x in row:
            x.border = BORDER
            if x.column > 1:
                x.alignment = Alignment(horizontal="center")
    for c in (1, 2):
        ws.cell(c, 1).font = Font(bold=True, size=11)
    ws.column_dimensions["A"].width = 30
    for c in range(2, len(head) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 8
    ws.column_dimensions["D"].width = 11


def open_file(path):
    os.startfile(path)  # Windows
